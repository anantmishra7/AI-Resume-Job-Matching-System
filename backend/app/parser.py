"""
Multi-format document parsing and section segmentation module.
Supports PDF (vector and messy layouts), DOCX, and TXT files.
"""

import io
import re
from typing import Dict, Any, List
from pypdf import PdfReader
from docx import Document

# Section demarcation headers
SECTION_PATTERNS = {
    "summary": r"(?:summary|objective|profile|about\s+me|professional\s+summary)",
    "experience": r"(?:experience|work\s+experience|employment\s+history|work\s+history|professional\s+experience)",
    "education": r"(?:education|academic\s+background|academic\s+history|qualifications)",
    "skills": r"(?:skills|technical\s+skills|core\s+competencies|technologies|areas\s+of\s+expertise|tools)",
    "projects": r"(?:projects|personal\s+projects|academic\s+projects|key\s+projects)",
    "certifications": r"(?:certifications|certificates|licenses|credentials)",
}

def clean_text(text: str) -> str:
    """Normalize whitespace, tabs, and standard resume bullet points."""
    text = re.sub(r'[\r\f\v]', '\n', text)
    text = re.sub(r'[\u2022\u2023\u25E6\u2043\u2219\u25AA\u25AB\u25CF\u25CB\u25C6\u25BA]', ' * ', text)
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def parse_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using pypdf with multi-page handling and layout normalization."""
    reader = PdfReader(io.BytesIO(file_bytes))
    extracted = []
    for page_idx, page in enumerate(reader.pages):
        try:
            page_text = page.extract_text() or ""
            if page_text.strip():
                extracted.append(page_text)
        except Exception:
            continue
    full_text = "\n\n".join(extracted)
    return clean_text(full_text)

def parse_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX including tables and paragraphs."""
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return clean_text("\n".join(paragraphs))

def parse_txt(file_bytes: bytes) -> str:
    """Extract text from raw text / markdown bytes with encoding fallback."""
    for enc in ['utf-8', 'latin-1', 'cp1252', 'utf-16']:
        try:
            return clean_text(file_bytes.decode(enc))
        except UnicodeDecodeError:
            continue
    return clean_text(file_bytes.decode('utf-8', errors='ignore'))

def parse_document(filename: str, file_bytes: bytes) -> str:
    """Route document to appropriate parser based on file extension."""
    lower_name = filename.lower()
    if lower_name.endswith('.pdf'):
        return parse_pdf(file_bytes)
    elif lower_name.endswith('.docx') or lower_name.endswith('.doc'):
        return parse_docx(file_bytes)
    else:
        return parse_txt(file_bytes)

def segment_sections(raw_text: str) -> Dict[str, str]:
    """
    Split resume text into semantic sections using regex boundary detection.
    Returns dictionary with keys: summary, experience, education, skills, projects, certifications, header.
    """
    lines = raw_text.splitlines()
    sections: Dict[str, List[str]] = {
        "header": [],
        "summary": [],
        "experience": [],
        "education": [],
        "skills": [],
        "projects": [],
        "certifications": [],
    }

    current_section = "header"

    # Pre-compile header regexes
    header_patterns = []
    for sec_name, pattern in SECTION_PATTERNS.items():
        rx = re.compile(rf"^(?:[#\*\-\s]{{0,3}})?({pattern})(?:[:\s\*\-#]{{0,3}})?$", re.IGNORECASE)
        header_patterns.append((sec_name, rx))

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        matched_new_section = False
        # Only check short lines (headings are rarely > 6 words)
        if len(stripped.split()) <= 6:
            for sec_name, rx in header_patterns:
                if rx.match(stripped):
                    current_section = sec_name
                    matched_new_section = True
                    break

        if not matched_new_section:
            sections[current_section].append(stripped)

    return {k: "\n".join(v).strip() for k, v in sections.items()}
